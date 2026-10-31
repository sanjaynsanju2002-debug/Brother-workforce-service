"""Payroll arithmetic, company isolation, saved snapshots and branded reports."""

import copy
import os
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import httpx
import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader
from pymongo.errors import DuplicateKeyError

os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "bws_payroll_test")

from lib import auth
from lib.invoices import InvoiceProfile
from lib.payroll import Payroll, PayrollInput, PayrollLine, calculate_payroll, sunday_count
from lib.payroll_exports import export_payroll_pdf, export_payroll_xlsx
from routers import assignments, payroll
from server import app
from tests.test_admin_session_and_application import FakeDB
from tests.test_finance import Collection as BaseCollection


def payload():
    return {"company_id": "company-a", "month": "2026-10", "notes": "Attendance approved by supervisor.", "lines": [
        {"assignment_id": "person-a", "worked_days": "26", "daily_rate": "500", "esi_percentage": "0.75", "pf_percentage": "12"},
        {"assignment_id": "person-b", "worked_days": "25.5", "daily_rate": "600", "esi_percentage": "0.75", "pf_percentage": "12"},
    ]}


def saved_payroll(count=2):
    data = payload()
    data["lines"] = [{**data["lines"][i % 2], "assignment_id": f"person-{i}"} for i in range(count)]
    value = PayrollInput(**data)
    amounts, totals = calculate_payroll(value)
    lines = [PayrollLine(**line.model_dump(), employee_id="0001" if i == 0 else f"BWS-{i:04}", full_name="Ravi Kumar" if i == 0 else f"Employee {i + 1}", **amounts[i].model_dump(exclude={"assignment_id"})) for i, line in enumerate(value.lines)]
    now = datetime.now(timezone.utc)
    return Payroll(id=payroll.register_id(value.company_id, value.month), **value.model_dump(exclude={"lines"}), lines=lines, company_name="Example Manufacturing Pvt Ltd", company_location="Hebbal Industrial Area, Mysore", supplier=InvoiceProfile(), totals=totals, revision=1, created_at=now, updated_at=now)


def test_gross_esi_pf_payable_and_fractional_days_use_decimal_arithmetic():
    rows, totals = calculate_payroll(PayrollInput(**payload()))
    assert (rows[0].gross_wages, rows[0].esi_deduction, rows[0].pf_deduction, rows[0].payable) == tuple(map(Decimal, ("13000.00", "97.50", "1560.00", "11342.50")))
    assert (rows[1].gross_wages, rows[1].esi_deduction, rows[1].pf_deduction, rows[1].payable) == tuple(map(Decimal, ("15300.00", "114.75", "1836.00", "13349.25")))
    assert totals.worked_days == Decimal("51.50")
    assert totals.employee_count == 2
    assert totals.gross_wages == Decimal("28300.00")
    assert totals.esi_deduction == Decimal("212.25")
    assert totals.pf_deduction == Decimal("3396.00")
    assert totals.payable == Decimal("24691.75")
    assert rows[0].sunday_holidays == rows[1].sunday_holidays == 0


def test_monthly_register_matches_salary_sheet_and_adds_sundays_from_selected_month():
    data = {
        "company_id": "company-a", "month": "2026-09", "lines": [{
            "assignment_id": "person-a", "calculation_method": "Monthly", "worked_days": "24",
            "base_salary": "15000", "days_absent": "2", "credit_leave": "0",
            "government_holidays": "1", "salary_advance": "0", "penalty": "0", "allowance": "0",
            "overtime_hours": "10", "overtime_rate": "56.60", "esi_percentage": "0.75",
            "esi_employer_percentage": "3.25", "pf_percentage": "0", "pf_employer_percentage": "0",
        }],
    }
    rows, totals = calculate_payroll(PayrollInput(**data))
    row = rows[0]
    assert sunday_count("2026-09") == row.sunday_holidays == 4
    assert row.loss_of_pay == Decimal("2.00")
    assert row.paid_days == Decimal("29.00")
    assert row.per_day_salary == Decimal("500.00")
    assert row.per_hour_salary == Decimal("55.56")
    assert row.leave_deduction == Decimal("1000.00")
    assert row.gross_wages == Decimal("14000.00")
    assert row.esi_deduction == Decimal("105.00")
    assert row.esi_employer_contribution == Decimal("455.00")
    assert row.overtime_amount == Decimal("566.00")
    assert row.payable == Decimal("14461.00")
    assert totals.paid_days == Decimal("29.00")


def test_monthly_register_uses_automatic_per_hour_rate_when_overtime_rate_is_zero():
    value = PayrollInput(**{
        "company_id": "company-a", "month": "2026-10", "lines": [{
            "assignment_id": "person-a", "calculation_method": "Monthly", "worked_days": "23",
            "base_salary": "27900", "days_absent": "1", "credit_leave": "1",
            "government_holidays": "1", "overtime_hours": "2", "overtime_rate": "0",
        }],
    })
    row = calculate_payroll(value)[0][0]
    assert row.sunday_holidays == 4
    assert row.loss_of_pay == 0
    assert row.per_day_salary == Decimal("900.00")
    assert row.effective_overtime_rate == Decimal("100.00")
    assert row.overtime_amount == Decimal("200.00")
    assert row.payable == Decimal("28100.00")


def test_each_deduction_rounds_to_paise_independently_and_zero_days_pay_zero():
    data = payload()
    data["lines"][0].update(worked_days="0.5", daily_rate="0.10", esi_percentage="10", pf_percentage="10")
    data["lines"][1].update(worked_days="0", esi_percentage="100", pf_percentage="0")
    rows, totals = calculate_payroll(PayrollInput(**data))
    assert rows[0].gross_wages == Decimal("0.05")
    assert rows[0].esi_deduction == rows[0].pf_deduction == Decimal("0.01")
    assert totals.payable == Decimal("0.03")
    assert rows[1].gross_wages == rows[1].payable == 0


@pytest.mark.parametrize("change", [{"worked_days": "-1"}, {"worked_days": "32"}, {"worked_days": "1.001"}, {"daily_rate": "NaN"}, {"daily_rate": "Infinity"}, {"daily_rate": "-5"}, {"daily_rate": "0.001"}, {"esi_percentage": "101"}, {"pf_percentage": "-1"}, {"pf_percentage": "NaN"}, {"pf_percentage": "12.001"}, {"esi_percentage": "60", "pf_percentage": "50"}, {"worked_days": "1", "daily_rate": "0.01", "esi_percentage": "50", "pf_percentage": "50"}])
def test_invalid_inputs_cannot_produce_negative_or_nonfinite_pay(change):
    data = payload()
    data["lines"][0].update(change)
    with pytest.raises(ValidationError):
        PayrollInput(**data)


def test_calendar_limits_and_duplicate_employees():
    data = payload()
    data.update(month="2026-02")
    data["lines"][0]["worked_days"] = "29"
    with pytest.raises(ValidationError):
        PayrollInput(**data)
    data["month"] = "2028-02"
    assert PayrollInput(**data).lines[0].worked_days == 29
    data["lines"][1]["assignment_id"] = data["lines"][0]["assignment_id"]
    with pytest.raises(ValidationError):
        PayrollInput(**data)
    with pytest.raises(ValidationError):
        PayrollInput(**{**payload(), "month": "2026-13"})


class Cursor:
    def __init__(self, rows, projection):
        self.rows, self.projection = copy.deepcopy(rows), projection

    def sort(self, key, direction):
        self.rows.sort(key=lambda row: row.get(key), reverse=direction == -1)
        return self

    def skip(self, count):
        self.rows = self.rows[count:]
        return self

    def limit(self, count):
        self.rows = self.rows[:count]
        return self

    async def to_list(self, count):
        return [BaseCollection.project(row, self.projection) for row in self.rows[:count]]


class Collection(BaseCollection):
    def find(self, query, projection=None):
        return Cursor([row for row in self.records if self.matches(row, query)], projection)


def fake_db(monkeypatch):
    fake = FakeDB()
    fake.clients, fake.assignments, fake.payroll = Collection(), Collection(), Collection()
    fake.clients.records = [{"id": "company-a", "name": "Factory A", "location": "Mysore", "visible": False}, {"id": "company-b", "name": "Factory B"}]
    fake.assignments.records = [
        {"id": "person-a", "client_id": "company-a", "full_name": "Ravi Kumar", "employee_id": "0001", "salary": "500", "salary_basis": "Daily", "status": "Active"},
        {"id": "person-b", "client_id": "company-a", "full_name": "Anitha", "salary": "15600", "salary_basis": "Monthly", "status": "Completed"},
        {"id": "person-foreign", "client_id": "company-b", "full_name": "Another employee", "salary": "750"},
    ]
    for module in (auth, assignments, payroll):
        monkeypatch.setattr(module, "db", fake)
    return fake


@pytest.mark.asyncio
async def test_admin_access_company_isolation_and_input_validation(monkeypatch):
    fake_db(monkeypatch)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        for path in ("", "/companies", "/draft?company_id=company-a&month=2026-10", "/missing", "/missing/download/pdf"):
            assert (await client.get("/api/admin/payroll" + path)).status_code == 401
        for path in ("", "/preview"):
            assert (await client.post("/api/admin/payroll" + path, json=payload())).status_code == 401
        assert (await client.post("/api/admin/login", json={"pin": "123456"})).status_code == 200
        assert (await client.post("/api/admin/payroll", json=payload(), headers={"Origin": "https://untrusted.invalid"})).status_code == 403
        companies = (await client.get("/api/admin/payroll/companies")).json()
        assert {company["id"] for company in companies} == {"company-a", "company-b"}
        draft = (await client.get("/api/admin/payroll/draft?company_id=company-a&month=2026-10")).json()
        assert draft["payroll"] is None
        assert {person["assignment_id"] for person in draft["employees"]} == {"person-a", "person-b"}
        assert next(person for person in draft["employees"] if person["assignment_id"] == "person-a")["salary_basis"] == "Daily"
        assert next(person for person in draft["employees"] if person["assignment_id"] == "person-b")["salary"] == "15600"
        assert (await client.get("/api/admin/payroll/draft?company_id=unknown&month=2026-10")).status_code == 404
        foreign = payload()
        foreign["lines"][0]["assignment_id"] = "person-foreign"
        assert (await client.post("/api/admin/payroll", json=foreign)).status_code == 422
        for lines in ([], [{**payload()["lines"][0], "worked_days": "0"}], [{**payload()["lines"][0], "daily_rate": "0"}]):
            assert (await client.post("/api/admin/payroll", json={**payload(), "lines": lines})).status_code == 422
        preview = (await client.post("/api/admin/payroll/preview", json=payload())).json()
        assert preview["totals"]["payable"] == "24691.75"


@pytest.mark.asyncio
async def test_monthly_register_saves_automatic_sundays_and_employer_contributions(monkeypatch):
    fake = fake_db(monkeypatch)
    fake.clients.records[0]["name"] = "Dhatu Organics & Naturals"
    body = {
        "company_id": "company-a", "month": "2026-09", "lines": [{
            "assignment_id": "person-b", "calculation_method": "Monthly", "worked_days": "21", "base_salary": "14400",
            "days_absent": "4", "credit_leave": "0", "government_holidays": "1", "salary_advance": "0",
            "penalty": "106", "allowance": "0", "overtime_hours": "9.26", "overtime_rate": "53.35",
            "esi_percentage": "0.75", "esi_employer_percentage": "3.25", "pf_percentage": "0", "pf_employer_percentage": "0",
        }],
    }
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        saved = await client.post("/api/admin/payroll", json=body)
        assert saved.status_code == 200, saved.text
        line = saved.json()["lines"][0]
        assert line["sunday_holidays"] == 4 and line["paid_days"] == "26.00"
        assert line["loss_of_pay"] == "4.00" and line["leave_deduction"] == "1920.00"
        assert line["gross_wages"] == "12374.00" and line["esi_employer_contribution"] == "402.16"


@pytest.mark.asyncio
async def test_other_companies_cannot_use_dhatu_monthly_register(monkeypatch):
    fake_db(monkeypatch)
    body = {
        "company_id": "company-a", "month": "2026-09", "lines": [{
            "assignment_id": "person-b", "calculation_method": "Monthly", "worked_days": "21",
            "base_salary": "14400", "days_absent": "4", "government_holidays": "1",
        }],
    }
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        for path in ("/api/admin/payroll/preview", "/api/admin/payroll"):
            response = await client.post(path, json=body)
            assert response.status_code == 422
            assert response.json()["detail"] == payroll.DHATU_MONTHLY_ONLY


@pytest.mark.asyncio
async def test_save_retry_stale_updates_history_and_downloads_keep_snapshots(monkeypatch):
    fake = fake_db(monkeypatch)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        created = await client.post("/api/admin/payroll", json=payload())
        assert created.status_code == 200, created.text
        saved = created.json()
        assert saved["totals"]["payable"] == "24691.75" and saved["revision"] == 1
        assert (await client.post("/api/admin/payroll", json=payload())).json() == saved
        assert len(fake.payroll.records) == 1
        assert (await client.post("/api/admin/payroll", json={**payload(), "notes": "changed"})).status_code == 409
        changed_data = {**payload(), "notes": "Updated attendance", "revision": 1, "created_at": saved["created_at"]}
        changed_data["lines"][0]["worked_days"] = "25"
        changed = await client.post("/api/admin/payroll", json=changed_data)
        assert changed.status_code == 200, changed.text
        changed = changed.json()
        assert changed["revision"] == 2
        assert changed["totals"]["payable"] == "24255.50"
        assert (await client.post("/api/admin/payroll", json=changed_data)).json()["revision"] == 2
        assert (await client.post("/api/admin/payroll", json={**payload(), "revision": 1})).status_code == 409
        path = "/api/admin/payroll/" + saved["id"]
        history = (await client.get("/api/admin/payroll?company_id=company-a")).json()
        assert history["total"] == 1
        assert "lines" not in history["items"][0] and "input_hash" not in history["items"][0]
        assert (await client.get("/api/admin/payroll?company_id=company-b")).json()["total"] == 0
        fake.clients.records = []
        fake.assignments.records = []
        assert (await client.get(path)).json() == changed
        draft = (await client.get("/api/admin/payroll/draft?company_id=company-a&month=2026-10")).json()
        assert len(draft["employees"]) == 2
        assert (await client.get("/api/admin/payroll/companies")).json()[0]["saved_only"] is True
        assert (await client.post("/api/admin/payroll", json={**changed_data, "revision": 2, "notes": "Archived company"})).json()["revision"] == 3
        for format, signature in (("pdf", b"%PDF"), ("xlsx", b"PK")):
            response = await client.get(f"{path}/download/{format}")
            assert response.status_code == 200, response.text[:200] if response.status_code != 200 else ""
            assert response.content.startswith(signature)
            assert response.headers["cache-control"] == "private, no-store"
            assert "Factory-A-2026-10" in response.headers["content-disposition"]
        assert (await client.get(f"{path}/download/csv")).status_code == 404
        assert (await client.get("/api/admin/payroll/missing/download/pdf")).status_code == 404


@pytest.mark.asyncio
async def test_competing_create_and_update_cannot_overwrite_another_save(monkeypatch):
    fake = fake_db(monkeypatch)
    insert = fake.payroll.insert_one

    async def competing_insert(record):
        await insert({**record, "notes": "Concurrent save", "input_hash": "other"})
        raise DuplicateKeyError("competing create")

    monkeypatch.setattr(fake.payroll, "insert_one", competing_insert)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        assert (await client.post("/api/admin/payroll", json=payload())).status_code == 409
        assert len(fake.payroll.records) == 1

        async def lost_update(query, changes):
            return SimpleNamespace(matched_count=0)

        monkeypatch.setattr(fake.payroll, "update_one", lost_update)
        assert (await client.post("/api/admin/payroll", json={**payload(), "revision": 1, "created_at": fake.payroll.records[0]["created_at"]})).status_code == 409
        assert fake.payroll.records[0]["notes"] == "Concurrent save"


@pytest.mark.asyncio
async def test_delete_is_private_version_checked_and_only_removes_the_selected_payroll(monkeypatch):
    fake = fake_db(monkeypatch)
    workforce = copy.deepcopy(fake.assignments.records)
    companies = copy.deepcopy(fake.clients.records)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        assert (await client.delete("/api/admin/payroll/missing", params={"revision": 1, "updated_at": datetime.now(timezone.utc).isoformat()})).status_code == 401
        await client.post("/api/admin/login", json={"pin": "123456"})
        target = (await client.post("/api/admin/payroll", json=payload())).json()
        other = (await client.post("/api/admin/payroll", json={**payload(), "month": "2026-09"})).json()
        path = f"/api/admin/payroll/{target['id']}"
        params = {"revision": target["revision"], "updated_at": target["updated_at"]}
        for invalid in ({}, {"revision": 1}, {"updated_at": target["updated_at"]}, {**params, "revision": 0}, {**params, "updated_at": "not-a-date"}, {**params, "updated_at": "2026-10-16T12:00:00"}):
            assert (await client.delete(path, params=invalid)).status_code == 422
        assert (await client.delete(path, params={**params, "revision": 2})).status_code == 409
        assert (await client.delete(path, params={**params, "updated_at": "2026-01-01T00:00:00Z"})).status_code == 409
        assert (await client.delete(path, params=params, headers={"Origin": "https://another.invalid"})).status_code == 403
        assert len(fake.payroll.records) == 2
        response = await client.delete(path, params=params)
        assert response.status_code == 204 and response.content == b""
        assert [row["id"] for row in fake.payroll.records] == [other["id"]]
        assert (await client.get(f"/api/admin/payroll/{other['id']}")).json() == other
        assert (await client.get("/api/admin/payroll?company_id=company-a")).json()["total"] == 1
        draft = (await client.get("/api/admin/payroll/draft?company_id=company-a&month=2026-10")).json()
        assert draft["payroll"] is None and len(draft["employees"]) == 2
        for suffix in ("", "/download/pdf", "/download/xlsx"):
            assert (await client.get(path + suffix)).status_code == 404
        assert (await client.delete(path, params=params)).status_code == 404
        assert fake.assignments.records == workforce
        assert fake.clients.records == companies


@pytest.mark.asyncio
async def test_deleted_month_can_be_recreated_without_stale_deletes_or_saves_overwriting_it(monkeypatch):
    fake = fake_db(monkeypatch)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        old = (await client.post("/api/admin/payroll", json=payload())).json()
        path = f"/api/admin/payroll/{old['id']}"
        params = {"revision": old["revision"], "updated_at": old["updated_at"]}
        assert (await client.delete(path, params=params)).status_code == 204
        replacement = (await client.post("/api/admin/payroll", json={**payload(), "notes": "Replacement payroll"})).json()
        assert replacement["id"] == old["id"] and replacement["revision"] == old["revision"] == 1
        assert replacement["created_at"] != old["created_at"]
        assert (await client.delete(path, params=params)).status_code == 409
        assert (await client.post("/api/admin/payroll", json={**payload(), "revision": old["revision"], "created_at": old["created_at"], "notes": "Stale edit"})).status_code == 409
        assert (await client.get(path)).json() == replacement
        changed = await client.post("/api/admin/payroll", json={**payload(), "revision": 1, "created_at": replacement["created_at"], "notes": "Approved replacement"})
        assert changed.status_code == 200, changed.text
        assert changed.json()["revision"] == 2
        assert (await client.delete(path, params={"revision": 1, "updated_at": replacement["updated_at"]})).status_code == 409
        assert len(fake.payroll.records) == 1


@pytest.mark.asyncio
async def test_deletion_and_save_match_the_reviewed_record_atomically(monkeypatch):
    fake = fake_db(monkeypatch)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        saved = (await client.post("/api/admin/payroll", json=payload())).json()
        delete = fake.payroll.delete_one

        async def competing_change(query):
            # The row changes after the endpoint's read but before Mongo deletes.
            fake.payroll.records[0]["updated_at"] = (datetime.now(timezone.utc) + timedelta(seconds=1)).isoformat()
            return await delete(query)

        monkeypatch.setattr(fake.payroll, "delete_one", competing_change)
        path = f"/api/admin/payroll/{saved['id']}"
        assert (await client.delete(path, params={"revision": 1, "updated_at": saved["updated_at"]})).status_code == 409
        assert len(fake.payroll.records) == 1
        update = fake.payroll.update_one

        async def competing_replacement(query, changes):
            # Deleting and recreating a month resets its revision to 1.
            fake.payroll.records[0]["created_at"] = (datetime.now(timezone.utc) + timedelta(seconds=1)).isoformat()
            return await update(query, changes)

        monkeypatch.setattr(fake.payroll, "update_one", competing_replacement)
        assert (await client.post("/api/admin/payroll", json={**payload(), "revision": 1, "created_at": saved["created_at"], "notes": "Stale save"})).status_code == 409
        assert fake.payroll.records[0]["revision"] == 1
        assert fake.payroll.records[0]["notes"] == payload()["notes"]


@pytest.mark.asyncio
async def test_workforce_salary_create_update_clear_and_legacy_defaults(monkeypatch):
    fake = fake_db(monkeypatch)
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        data = {"client_id": "company-a", "full_name": "New Employee", "salary": "18500.50", "salary_basis": "Monthly"}
        response = await client.post("/api/admin/assignments", json=data)
        assert response.status_code == 200, response.text
        person = response.json()
        assert person["salary"] == "18500.50"
        stored = next(row for row in fake.assignments.records if row["id"] == person["id"])
        assert isinstance(stored["salary"], str) and isinstance(stored["created_at"], datetime)
        path = "/api/admin/assignments/" + person["id"]
        changed = (await client.patch(path, json={"salary": "600", "salary_basis": "Daily"})).json()
        assert changed["salary"] == "600" and changed["salary_basis"] == "Daily"
        assert (await client.patch(path, json={"salary": None})).json()["salary"] is None
        for change in ({"salary": "NaN"}, {"salary": "-1"}, {"salary": "0.001"}, {"salary_basis": None}, {"salary_basis": "Weekly"}):
            assert (await client.patch(path, json=change)).status_code == 422
        draft = (await client.get("/api/admin/payroll/draft?company_id=company-b&month=2026-10")).json()
        assert draft["employees"][0]["salary_basis"] == "Monthly"


@pytest.mark.parametrize("count", [1, 2, 40])
def test_excel_preserves_logo_print_settings_text_ids_and_dynamic_formulas(count):
    saved = saved_payroll(count)
    saved.lines[0].full_name = '=HYPERLINK("https://example.invalid","Name")'
    with ZipFile(BytesIO(export_payroll_xlsx(saved))) as archive:
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        cells = {cell.get("r"): cell for cell in root.findall(".//s:c", ns)}
        assert any(name.startswith("xl/media/") for name in archive.namelist())
        assert cells["B16"].get("t") == cells["C16"].get("t") == "inlineStr"
        assert cells["B16"].find("s:is/s:t", ns).text == "0001"
        assert cells["C16"].find("s:is/s:t", ns).text.startswith("=HYPERLINK")
        assert cells["C16"].find("s:f", ns) is None
        assert cells["A9"].find("s:is/s:t", ns).text == "Company: " + saved.company_name
        total_row = 16 + count
        for i, line in enumerate(saved.lines, 16):
            assert cells[f"F{i}"].find("s:f", ns).text == f"ROUND(D{i}*E{i},2)"
            assert cells[f"H{i}"].find("s:f", ns).text == f"ROUND(F{i}*G{i}/100,2)"
            assert cells[f"J{i}"].find("s:f", ns).text == f"ROUND(F{i}*I{i}/100,2)"
            assert cells[f"K{i}"].find("s:f", ns).text == f"F{i}-H{i}-J{i}"
            assert Decimal(cells[f"K{i}"].find("s:v", ns).text) == line.payable
        for col, value in {"F": saved.totals.gross_wages, "H": saved.totals.esi_deduction, "J": saved.totals.pf_deduction, "K": saved.totals.payable}.items():
            assert cells[f"{col}{total_row}"].find("s:f", ns).text == f"SUM({col}16:{col}{total_row - 1})"
            assert Decimal(cells[f"{col}{total_row}"].find("s:v", ns).text) == value
        assert cells["J13"].find("s:f", ns).text == f"K{total_row}"
        assert Decimal(cells["J13"].find("s:v", ns).text) == saved.totals.payable
        assert cells[f"A{total_row + 2}"].find("s:is/s:t", ns).text == "Notes: " + saved.notes
        assert f"A{total_row}:C{total_row}" in [merge.get("ref") for merge in root.findall("s:mergeCells/s:mergeCell", ns)]
        assert root.find("s:autoFilter", ns).get("ref") == f"A15:K{total_row - 1}"
        assert root.find("s:pageSetup", ns).get("orientation") == "landscape"
        assert root.find("s:dataValidations", ns).get("count") == "4"
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        assert workbook.find("s:definedNames/s:definedName[@name='_xlnm.Print_Area']", ns).text.endswith(f"$K${count + 19}")
        assert "Employee name" not in " ".join(cell.find("s:is/s:t", ns).text or "" for cell in cells.values() if cell.find("s:is/s:t", ns) is not None and cell.get("r") != "C15")


def test_monthly_register_download_contains_sheet_columns_sundays_and_formulas():
    value = PayrollInput(**{
        "company_id": "company-a", "month": "2026-09", "lines": [{
            "assignment_id": "person-a", "calculation_method": "Monthly", "worked_days": "24", "base_salary": "15000",
            "days_absent": "2", "credit_leave": "0", "government_holidays": "1", "overtime_hours": "10",
            "overtime_rate": "56.60", "esi_percentage": "0.75", "esi_employer_percentage": "3.25",
            "pf_percentage": "0", "pf_employer_percentage": "0",
        }],
    })
    amounts, totals = calculate_payroll(value)
    line = PayrollLine(**value.lines[0].model_dump(), employee_id="0001", full_name="Shivaprasad", start_date="2022-07-27", **amounts[0].model_dump(exclude={"assignment_id"}))
    now = datetime.now(timezone.utc)
    saved = Payroll(id="monthly", **value.model_dump(exclude={"lines"}), lines=[line], company_name="DONPL", company_location="Mysore", supplier=InvoiceProfile(), totals=totals, revision=1, created_at=now, updated_at=now)
    book = load_workbook(BytesIO(export_payroll_xlsx(saved)), data_only=False)
    sheet = book["Payroll"]
    assert sheet["I8"].value == "Sunday Holiday"
    assert sheet["I9"].value == 4
    assert sheet["G9"].value == "=MAX(0,E9-F9)"
    assert sheet["K9"].value == "=F9+H9+I9+J9"
    assert sheet["AC9"].value == "=ROUND(Q9-S9-W9+AB9,2)"
    assert sheet["AE9"].value == "=ROUND(L9/30,2)"
    assert sheet.freeze_panes == "E9"
    text = " ".join(page.extract_text() for page in PdfReader(BytesIO(export_payroll_pdf(saved))).pages)
    for expected in ("SALARY REGISTER", "Shivaprasad", "Sunday", "14,461.00"):
        assert expected in text


def test_pdf_contains_branding_separate_deductions_and_repeated_headers_on_later_pages():
    single = saved_payroll()
    reader = PdfReader(BytesIO(export_payroll_pdf(single)))
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    for expected in (single.company_name, single.supplier.name, "Boregowda", "October 2026", "PF deduction", "212.25", "3,396.00", "24,691.75", "0001", "Ravi Kumar"):
        assert expected in text
    assert len(reader.pages[0].images) > 0
    multi = PdfReader(BytesIO(export_payroll_pdf(saved_payroll(55))))
    assert len(multi.pages) > 1
    for page in multi.pages:
        assert "Employee ID" in page.extract_text()
        assert "PF deduction" in page.extract_text()
    assert "Employee 55" in " ".join(page.extract_text() for page in multi.pages)


def paid_days_input(month="2026-10", **changes):
    line = {"assignment_id": "person-a", "calculation_method": "Monthly", "monthly_pay_basis": "PaidDays", "base_salary": "15000", "hours_per_day": "8", "worked_days": "25", "paid_sundays": 4, "days_absent": "2", "esi_percentage": "0.75"}
    line.update(changes)
    return PayrollInput(company_id="company-a", month=month, lines=[line])


def test_worked_days_drive_gross_with_sundays_and_esi_then_net_adjustments():
    value = paid_days_input(base_salary="15500", allowance="100", overtime_hours="2", salary_advance="500", penalty="25", esi_employer_percentage="3.25")
    row = calculate_payroll(value)[0][0]
    assert row.paid_days == Decimal("29")
    assert row.loss_of_pay == Decimal("2")
    assert row.per_day_salary == Decimal("500")
    assert row.per_hour_salary == Decimal("62.50")
    assert row.overtime_amount == Decimal("125")
    assert row.gross_wages == Decimal("14725")  # 29 x 500 + 100 + 125
    assert row.esi_deduction == Decimal("110.44")
    assert row.payable == Decimal("14089.56")  # gross - ESI - advance - penalty
    assert row.esi_employer_contribution == Decimal("478.56")
    fewer_days = calculate_payroll(paid_days_input(base_salary="15500", worked_days="23", days_absent="4"))[0][0]
    assert fewer_days.gross_wages == Decimal("13500")


@pytest.mark.parametrize("month,worked,sundays", [("2026-10", "27", 4), ("2026-09", "26", 4), ("2026-02", "24", 4), ("2028-02", "25", 4), ("2026-05", "26", 5)])
def test_full_month_pays_exact_ctc_without_rounding_drift(month, worked, sundays):
    row = calculate_payroll(paid_days_input(month, worked_days=worked, paid_sundays=sundays, days_absent="0"))[0][0]
    assert row.gross_wages == Decimal("15000")
    assert row.esi_deduction == Decimal("112.50")
    assert row.payable == Decimal("14887.50")


def test_paid_leave_half_days_and_sunday_override_are_counted_once():
    value = paid_days_input(base_salary="15500", worked_days="23.5", credit_leave="1.5", government_holidays="1", paid_sundays=3)
    row = calculate_payroll(value)[0][0]
    assert row.paid_days == Decimal("29") and row.loss_of_pay == Decimal("2")
    assert row.gross_wages == Decimal("14500")
    assert row.payable == Decimal("14391.25")


@pytest.mark.parametrize("changes", [{"worked_days": "28"}, {"paid_sundays": 5}, {"paid_sundays": "1.5"}, {"hours_per_day": "0"}, {"hours_per_day": "25"}, {"salary_advance": "15000"}, {"esi_percentage": "100", "salary_advance": "1"}])
def test_paid_days_reject_overlap_invalid_hours_sundays_and_excess_deductions(changes):
    with pytest.raises(ValidationError):
        paid_days_input(**changes)


def test_paid_days_exports_cached_values_and_recalculating_formulas_agree():
    value = paid_days_input(base_salary="15500", salary_advance="500", penalty="25")
    amounts, totals = calculate_payroll(value)
    line = PayrollLine(**value.lines[0].model_dump(), full_name="Payroll Example", **amounts[0].model_dump(exclude={"assignment_id"}))
    now = datetime.now(timezone.utc)
    saved = Payroll(id="paid-days", **value.model_dump(exclude={"lines"}), lines=[line], company_name="DONPL", company_location="Mysore", supplier=InvoiceProfile(), totals=totals, revision=1, created_at=now, updated_at=now)
    exported = export_payroll_xlsx(saved)
    sheet = load_workbook(BytesIO(exported), data_only=False).active
    assert sheet["Q9"].value == "=ROUND(L9*K9/31+P9+AB9,2)"
    assert sheet["AC9"].value == "=ROUND(Q9-S9-W9-M9-O9,2)"
    assert sheet["AF9"].value == 8
    cached = load_workbook(BytesIO(exported), data_only=True).active
    assert cached["Q9"].value == float(line.gross_wages) == 14500
    assert cached["S9"].value == float(line.esi_deduction) == 108.75
    assert cached["AC9"].value == float(line.payable) == 13866.25
    text = " ".join(page.extract_text() for page in PdfReader(BytesIO(export_payroll_pdf(saved))).pages)
    assert "13,866.25" in text and "hours/day 8" in text


@pytest.mark.asyncio
async def test_paid_days_manual_values_save_reopen_and_revise(monkeypatch):
    fake = fake_db(monkeypatch)
    fake.clients.records[0]["name"] = "DONPL"
    body = paid_days_input(base_salary="15500").model_dump(mode="json")
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="https://brothersworkforcesolutions.in") as client:
        await client.post("/api/admin/login", json={"pin": "123456"})
        response = await client.post("/api/admin/payroll", json=body)
        assert response.status_code == 200, response.text
        saved = response.json()
        draft = (await client.get("/api/admin/payroll/draft?company_id=company-a&month=2026-10")).json()
        assert draft["payroll"]["lines"][0]["hours_per_day"] == "8"
        assert draft["payroll"]["lines"][0]["paid_sundays"] == 4
        body.update(revision=saved["revision"], created_at=saved["created_at"])
        body["lines"][0].update(worked_days="24", days_absent="3", hours_per_day="7", esi_percentage="1")
        revised = await client.post("/api/admin/payroll", json=body)
        assert revised.status_code == 200, revised.text
        assert revised.json()["lines"][0]["gross_wages"] == "14000.00"
        assert revised.json()["lines"][0]["payable"] == "13860.00"
