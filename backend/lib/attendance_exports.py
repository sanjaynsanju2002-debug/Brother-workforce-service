"""Professional monthly and weekly attendance report downloads."""

import calendar
import csv
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

from lib.invoices import InvoiceProfile

ASSETS = Path(__file__).resolve().parent.parent / "assets"
STATUS_SHORT = {
    "Present": "P", "Absent": "A", "Half-day": "HD", "Paid leave": "PL",
    "Unpaid leave": "UL", "Leave": "UL", "Holiday": "H", "Weekly off": "WO",
    "Unmarked": "-",
}
PAID_VALUE = {
    "Present": Decimal("1"), "Half-day": Decimal("0.5"), "Paid leave": Decimal("1"),
    "Holiday": Decimal("1"), "Weekly off": Decimal("1"),
}
SHIFT_SHORT = {"General": "G", "Morning": "M", "Afternoon": "AF", "Night": "N"}


def month_label(month: str) -> str:
    return f"{calendar.month_name[int(month[5:])]} {month[:4]}"


def register_rows(register: dict):
    year, number = int(register["month"][:4]), int(register["month"][5:])
    days = calendar.monthrange(year, number)[1]
    sheets = {sheet["attendance_date"]: sheet for sheet in register["sheets"] if sheet.get("status") == "Approved"}
    lookups = {day: {row["assignment_id"]: row for row in sheet.get("rows", [])} for day, sheet in sheets.items()}
    rows = []
    for employee in register["employees"]:
        statuses, counts, payable = [], {}, Decimal("0")
        for day in range(1, days + 1):
            key = f"{register['month']}-{day:02d}"
            status = lookups.get(key, {}).get(employee["assignment_id"], {}).get("status", "Unmarked")
            statuses.append(STATUS_SHORT.get(status, "-"))
            if status != "Unmarked":
                counts[status] = counts.get(status, 0) + 1
            payable += PAID_VALUE.get(status, Decimal("0"))
        rows.append({**employee, "statuses": statuses, "counts": counts, "payable_days": payable})
    return days, rows, len(sheets)


def export_attendance_xlsx(register: dict, supplier: InvoiceProfile) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    days, rows, approved = register_rows(register)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Attendance"
    navy, orange, pale, border = "0F2444", "EA580C", "F1F5F9", "CBD5E1"
    last_column = 3 + days + 8
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet["A1"] = supplier.name
    sheet["A1"].font = Font(bold=True, size=16, color=navy)
    sheet["A2"] = supplier.address
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet["A3"] = " | ".join(filter(None, (supplier.phone, supplier.email, supplier.website)))
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    sheet["A5"] = "MONTHLY ATTENDANCE REGISTER"
    sheet["A5"].font = Font(bold=True, size=14, color=orange)
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=last_column)
    sheet["A6"] = f"Company: {register['company_name']}"
    sheet.merge_cells(start_row=6, start_column=1, end_row=6, end_column=3 + days)
    sheet.cell(6, 4 + days, f"Month: {month_label(register['month'])} | Approved days: {approved}")
    sheet.merge_cells(start_row=6, start_column=4 + days, end_row=6, end_column=last_column)
    headers = ["No.", "Employee ID", "Employee name", *range(1, days + 1), "P", "A", "HD", "PL", "UL", "H", "WO", "Payable days"]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(8, column, value)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(bold=True, color="FFFFFF", size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90 if isinstance(value, int) else 0)
    status_columns = {name: 4 + days + index for index, name in enumerate(("P", "A", "HD", "PL", "UL", "H", "WO"))}
    thin = Side(style="thin", color=border)
    for index, row in enumerate(rows, 1):
        excel_row = 8 + index
        values = [index, row.get("employee_id") or "", row["full_name"], *row["statuses"]]
        for column, value in enumerate(values, 1):
            sheet.cell(excel_row, column, value)
        normalized_counts = {STATUS_SHORT.get(status, status): count for status, count in row["counts"].items()}
        for label, column in status_columns.items():
            sheet.cell(excel_row, column, normalized_counts.get(label, 0))
        sheet.cell(excel_row, last_column, float(row["payable_days"]))
        for column in range(1, last_column + 1):
            cell = sheet.cell(excel_row, column)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="left" if column == 3 else "center", vertical="center")
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    sheet.freeze_panes = "D9"
    sheet.auto_filter.ref = f"A8:{get_column_letter(last_column)}{8 + len(rows)}"
    sheet.column_dimensions["A"].width = 6
    sheet.column_dimensions["B"].width = 15
    sheet.column_dimensions["C"].width = 28
    for column in range(4, 4 + days):
        sheet.column_dimensions[get_column_letter(column)].width = 4
    for column in range(4 + days, last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 10 if column == last_column else 6
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:8"
    sheet.print_area = f"A1:{get_column_letter(last_column)}{8 + len(rows)}"
    sheet.oddFooter.center.text = f"{supplier.name} | {month_label(register['month'])} | Page &P of &N"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_attendance_pdf(register: dict, supplier: InvoiceProfile) -> bytes:
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("AttendanceSans", "Vera.ttf"), ("AttendanceBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    days, rows, approved = register_rows(register)
    output = BytesIO()
    page = landscape(A3)
    doc = SimpleDocTemplate(output, pagesize=page, leftMargin=22, rightMargin=22, topMargin=28, bottomMargin=28,
                            title=f"Attendance - {register['company_name']} - {register['month']}", author=supplier.name)
    navy, orange, muted = (colors.HexColor(value) for value in ("#0F2444", "#EA580C", "#64748B"))
    body = ParagraphStyle("AttendanceBody", fontName="AttendanceSans", fontSize=6, leading=8, textColor=navy)
    bold = ParagraphStyle("AttendanceBold", parent=body, fontName="AttendanceBold")
    heading = ParagraphStyle("AttendanceHeading", parent=bold, fontSize=13, leading=16)
    small = ParagraphStyle("AttendanceSmall", parent=body, fontSize=5.5, leading=7, textColor=muted)

    def p(value, style=body):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)

    with ZipFile(ASSETS / "invoice-template.xlsx") as archive:
        logo = Image(BytesIO(archive.read("xl/media/image.png")), width=80, height=54.5)
    brand = [p(supplier.name, heading), p(supplier.address), p(" | ".join(filter(None, (supplier.phone, supplier.email, supplier.website))), small)]
    header = Table([[logo, brand]], colWidths=[95, doc.width - 95])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story = [header, Spacer(1, 4), p("MONTHLY ATTENDANCE REGISTER", heading), Spacer(1, 5),
             p(f"Company: {register['company_name']}   |   Month: {month_label(register['month'])}   |   Approved days: {approved}"), Spacer(1, 8)]
    headers = ["No.", "ID", "Employee", *[str(day) for day in range(1, days + 1)], "P", "A", "HD", "PL", "UL", "H", "WO", "Pay"]
    data = [[p(value, bold) for value in headers]]
    for index, row in enumerate(rows, 1):
        normalized_counts = {STATUS_SHORT.get(status, status): count for status, count in row["counts"].items()}
        values = [index, row.get("employee_id") or "-", row["full_name"], *row["statuses"],
                  *[normalized_counts.get(label, 0) for label in ("P", "A", "HD", "PL", "UL", "H", "WO")], f"{row['payable_days']:g}"]
        data.append([p(value, bold if column in {2, len(values) - 1} else body) for column, value in enumerate(values)])
    fixed = 24 + 54 + 112 + 8 * 28
    day_width = max(14, (doc.width - fixed) / days)
    widths = [24, 54, 112, *([day_width] * days), *([28] * 8)]
    table = LongTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (3, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.extend([table, Spacer(1, 7), p("Legend: P Present | A Absent | HD Half day | PL Paid leave | UL Unpaid leave | H Government holiday | WO Weekly off. Payable days include P, PL, H and WO as 1 day and HD as 0.5 day.", small)])

    def footer(canvas, document):
        canvas.saveState(); canvas.setFont("AttendanceSans", 6); canvas.setFillColor(muted)
        canvas.drawString(22, 14, f"{supplier.name[:70]} | Attendance {register['month']}")
        canvas.drawRightString(page[0] - 22, 14, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _weekly_day_label(value: str) -> str:
    day = date.fromisoformat(value)
    return day.strftime("%a %d %b")


def _weekly_status_cell(day: dict) -> str:
    status = STATUS_SHORT.get(day.get("status", "Unmarked"), "-")
    if status == "-":
        return status
    return f"{status}/{SHIFT_SHORT.get(day.get('shift'), 'G')}"


def _safe_spreadsheet(value) -> str:
    text = str(value or "")
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def export_attendance_weekly_csv(report: dict, supplier: InvoiceProfile) -> bytes:
    """Return an Excel-friendly UTF-8 CSV with one row per employee."""
    output = StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    status_names = ("Present", "Absent", "Half-day", "Paid leave", "Unpaid leave", "Holiday", "Weekly off")
    writer.writerow([
        "Organisation", "Company", "Week Start", "Week End", "Employee ID", "Employee Name", "Role",
        *[_weekly_day_label(day["date"]) for day in report["calendar_days"]],
        "Present", "Absent", "Half-day", "Paid Leave", "Unpaid Leave", "Holiday", "Weekly Off",
        "Payable Days", "Overtime Hours",
    ])
    for employee in report["employees"]:
        writer.writerow([
            _safe_spreadsheet(supplier.name), _safe_spreadsheet(report["company_name"]),
            report["week_start"], report["week_end"], _safe_spreadsheet(employee.get("employee_id")),
            _safe_spreadsheet(employee["full_name"]), _safe_spreadsheet(employee.get("role")),
            *[_weekly_status_cell(day) for day in employee["days"]],
            *[employee["status_counts"].get(status, 0) for status in status_names],
            employee["payable_days"], employee["overtime_hours"],
        ])
    return output.getvalue().encode("utf-8")


def export_attendance_weekly_xlsx(report: dict, supplier: InvoiceProfile) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Weekly Report"

    def write_text(cell, value):
        """Force user-entered text to remain text without displaying a quote prefix."""
        cell.value = str(value or "")
        cell.data_type = "s"

    navy, orange, border = "0F2444", "EA580C", "CBD5E1"
    status_names = ("Present", "Absent", "Half-day", "Paid leave", "Unpaid leave", "Holiday", "Weekly off")
    headers = [
        "No.", "Employee ID", "Employee name", "Role",
        *[f"{_weekly_day_label(day['date'])}\n{day['status']}" for day in report["calendar_days"]],
        "P", "A", "HD", "PL", "UL", "H", "WO", "Payable days", "OT hours",
    ]
    last_column = len(headers)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    write_text(sheet["A1"], supplier.name)
    sheet["A1"].font = Font(bold=True, size=16, color=navy)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    write_text(sheet["A2"], supplier.address)
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_column)
    write_text(sheet["A3"], " | ".join(filter(None, (supplier.phone, supplier.email, supplier.website))))
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=last_column)
    sheet["A5"] = "WEEKLY ATTENDANCE REPORT"
    sheet["A5"].font = Font(bold=True, size=14, color=orange)
    sheet.merge_cells(start_row=6, start_column=1, end_row=6, end_column=11)
    write_text(sheet["A6"], f"Company: {report['company_name']}")
    sheet.merge_cells(start_row=6, start_column=12, end_row=6, end_column=last_column)
    sheet.cell(6, 12, f"Week: {report['week_start']} to {report['week_end']} | Approved: {report['approved_days']} of 7")

    for column, value in enumerate(headers, 1):
        cell = sheet.cell(8, column, value)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(bold=True, color="FFFFFF", size=8)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[8].height = 34
    thin = Side(style="thin", color=border)
    status_start = 12
    for index, employee in enumerate(report["employees"], 1):
        excel_row = 8 + index
        values = [
            index, employee.get("employee_id"), employee["full_name"], employee.get("role"),
            *[_weekly_status_cell(day) for day in employee["days"]],
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(excel_row, column)
            if isinstance(value, str) or value is None:
                write_text(cell, value)
            else:
                cell.value = value
        for offset, status in enumerate(status_names):
            sheet.cell(excel_row, status_start + offset, employee["status_counts"].get(status, 0))
        sheet.cell(excel_row, last_column - 1, float(Decimal(employee["payable_days"])))
        sheet.cell(excel_row, last_column, float(Decimal(employee["overtime_hours"])))
        for column in range(1, last_column + 1):
            cell = sheet.cell(excel_row, column)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="left" if column in {3, 4} else "center", vertical="center")
            if index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")

    sheet.freeze_panes = "E9"
    sheet.auto_filter.ref = f"A8:{get_column_letter(last_column)}{8 + len(report['employees'])}"
    for column, width in {"A": 6, "B": 15, "C": 27, "D": 20}.items():
        sheet.column_dimensions[column].width = width
    for column in range(5, 12):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    for column in range(12, last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 12 if column >= last_column - 1 else 6
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:8"
    sheet.print_area = f"A1:{get_column_letter(last_column)}{8 + len(report['employees'])}"
    sheet.oddFooter.center.text = f"{supplier.name} | Weekly attendance | Page &P of &N"

    notes = [(day["date"], day["status"], day["notes"]) for day in report["calendar_days"] if day.get("notes")]
    if notes:
        notes_sheet = workbook.create_sheet("Daily Notes")
        notes_sheet.append(["Date", "Sheet status", "Notes"])
        for row in notes:
            notes_sheet.append([date.fromisoformat(row[0]), row[1], ""])
            notes_sheet.cell(notes_sheet.max_row, 1).number_format = "dd mmm yyyy"
            write_text(notes_sheet.cell(notes_sheet.max_row, 3), row[2])
        for cell in notes_sheet[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(bold=True, color="FFFFFF")
        notes_sheet.column_dimensions["A"].width = 18
        notes_sheet.column_dimensions["B"].width = 16
        notes_sheet.column_dimensions["C"].width = 80
        notes_sheet.sheet_view.showGridLines = False

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_attendance_weekly_pdf(report: dict, supplier: InvoiceProfile) -> bytes:
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("WeeklySans", "Vera.ttf"), ("WeeklyBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    output = BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(output, pagesize=page, leftMargin=20, rightMargin=20, topMargin=24, bottomMargin=26,
                            title=f"Weekly Attendance - {report['company_name']}", author=supplier.name)
    navy, orange, muted = (colors.HexColor(value) for value in ("#0F2444", "#EA580C", "#64748B"))
    body = ParagraphStyle("WeeklyBody", fontName="WeeklySans", fontSize=5.5, leading=7, textColor=navy)
    bold = ParagraphStyle("WeeklyBold", parent=body, fontName="WeeklyBold")
    heading = ParagraphStyle("WeeklyHeading", parent=bold, fontSize=13, leading=16)
    small = ParagraphStyle("WeeklySmall", parent=body, fontSize=5, leading=6.5, textColor=muted)
    table_header = ParagraphStyle("WeeklyTableHeader", parent=bold, textColor=colors.white)

    def p(value, style=body):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)

    with ZipFile(ASSETS / "invoice-template.xlsx") as archive:
        logo = Image(BytesIO(archive.read("xl/media/image.png")), width=72, height=49)
    brand = [p(supplier.name, heading), p(supplier.address), p(" | ".join(filter(None, (supplier.phone, supplier.email, supplier.website))), small)]
    header = Table([[logo, brand]], colWidths=[84, doc.width - 84])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story = [
        header, Spacer(1, 3), p("WEEKLY ATTENDANCE REPORT", heading), Spacer(1, 4),
        p(f"Company: {report['company_name']}   |   Week: {report['week_start']} to {report['week_end']}   |   Approved days: {report['approved_days']} of 7"),
        Spacer(1, 7),
    ]
    summary_labels = ("Employees", "Approved days", "Present", "Absent", "Half days", "Payable days", "OT hours")
    summary_values = (
        report["totals"]["employees"], report["approved_days"], report["status_counts"].get("Present", 0),
        report["status_counts"].get("Absent", 0), report["status_counts"].get("Half-day", 0),
        report["totals"]["payable_days"], report["totals"]["overtime_hours"],
    )
    summary = Table([
        [p(label, small) for label in summary_labels],
        [p(value, bold) for value in summary_values],
    ], colWidths=[doc.width / len(summary_labels)] * len(summary_labels))
    summary.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), .4, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), .25, colors.HexColor("#E2E8F0")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.extend([summary, Spacer(1, 7)])

    status_names = ("Present", "Absent", "Half-day", "Paid leave", "Unpaid leave", "Holiday", "Weekly off")
    headers = [
        "No.", "ID", "Employee", "Role",
        *[f"{_weekly_day_label(day['date'])}\n{day['status']}" for day in report["calendar_days"]],
        "P", "A", "HD", "PL", "UL", "H", "WO", "Pay", "OT",
    ]
    data = [[p(value, table_header) for value in headers]]
    for index, employee in enumerate(report["employees"], 1):
        values = [
            index, employee.get("employee_id") or "-", employee["full_name"], employee.get("role") or "-",
            *[_weekly_status_cell(day) for day in employee["days"]],
            *[employee["status_counts"].get(status, 0) for status in status_names],
            employee["payable_days"], employee["overtime_hours"],
        ]
        data.append([p(value, bold if column in {2, len(values) - 2} else body) for column, value in enumerate(values)])
    widths = [22, 48, 120, 70, *([42] * 7), *([22] * 7), 40, 40]
    table = LongTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (4, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)
    notes = [day for day in report["calendar_days"] if day.get("notes")]
    if notes:
        story.extend([Spacer(1, 7), p("DAILY NOTES", bold)])
        story.extend(p(f"{_weekly_day_label(day['date'])} ({day['status']}): {day['notes']}", small) for day in notes)
    story.extend([Spacer(1, 6), p(
        "Legend: P Present | A Absent | HD Half day | PL Paid leave | UL Unpaid leave | H Government holiday | WO Weekly off. "
        "Shift: G General | M Morning | AF Afternoon | N Night. Totals include approved days only.", small,
    )])

    def footer(canvas, document):
        canvas.saveState(); canvas.setFont("WeeklySans", 6); canvas.setFillColor(muted)
        canvas.drawString(20, 13, f"{supplier.name[:65]} | Weekly attendance {report['week_start']}")
        canvas.drawRightString(page[0] - 20, 13, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
