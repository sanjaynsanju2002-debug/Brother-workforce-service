"""Branded payroll downloads from the saved register, including ESI and PF."""

import calendar
import re
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from lib.payroll import Payroll

ASSETS = Path(__file__).resolve().parent.parent / "assets"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def month_label(month: str) -> str:
    return f"{calendar.month_name[int(month[5:])]} {month[:4]}"


def export_monthly_register_xlsx(payroll: Payroll) -> bytes:
    """Build the expanded monthly salary register used by the admin payroll screen."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payroll"
    sheet.sheet_view.showGridLines = False
    supplier = payroll.supplier
    last_column = 31
    last_letter = get_column_letter(last_column)
    navy, orange, pale, blue = "0F2444", "C2410C", "E2E8F0", "DBEAFE"
    thin = Side(style="thin", color=pale)

    sheet.merge_cells(f"A1:{last_letter}1")
    sheet["A1"] = supplier.name
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=navy)
    sheet["A1"].alignment = Alignment(horizontal="left")
    sheet.merge_cells(f"A2:{last_letter}2")
    sheet["A2"] = supplier.address
    sheet.merge_cells(f"A3:{last_letter}3")
    sheet["A3"] = " | ".join(filter(None, (supplier.phone, supplier.email, supplier.website, f"GSTIN: {supplier.gstin}" if supplier.gstin else "")))
    sheet.merge_cells(f"A5:{last_letter}5")
    sheet["A5"] = f"SALARY REGISTER · {payroll.company_name} · {month_label(payroll.month)}"
    sheet["A5"].font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    sheet["A5"].fill = PatternFill("solid", fgColor=navy)
    sheet["A5"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[5].height = 25
    sheet.merge_cells(f"A6:{last_letter}6")
    sheet["A6"] = f"Worksite: {payroll.company_location or 'Not supplied'} | Sundays: {payroll.lines[0].sunday_holidays if payroll.lines else 0} | Employees: {len(payroll.lines)}"
    sheet["A6"].font = Font(name="Arial", size=9, color="475569")

    headers = [
        "Sl. No.", "Employee ID", "Name", "Date of Joining", "Days Absent", "Credit Leave", "Loss of Pay (LOP)",
        "Govt Holiday", "Sunday Holiday", "Days Present", "Paid Days", "Salary", "Salary Loan / Advance",
        "Leave Deduction", "Penalty", "Allowance", "Salary Gross Amount", "Employee ESI %", "Employee ESI",
        "Employer ESI %", "Employer ESI", "Employee PF %", "Employee PF", "Employer PF %", "Employer PF",
        "OT / Extra Hours", "OT Hourly Rate", "OT / Extra Amount", "Salary Net Amount", "Per Hour Salary", "Per Day Salary",
    ]
    header_row = 8
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(header_row, column, label)
        cell.font = Font(name="Arial", size=8, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.row_dimensions[header_row].height = 48

    start = header_row + 1
    days = calendar.monthrange(int(payroll.month[:4]), int(payroll.month[5:]))[1]
    for index, line in enumerate(payroll.lines, start):
        monthly = line.calculation_method == "Monthly"
        values = [
            index - header_row, line.employee_id or "", line.full_name, line.start_date or "", line.days_absent if monthly else "",
            line.credit_leave if monthly else "", line.loss_of_pay if monthly else "", line.government_holidays if monthly else "",
            line.sunday_holidays if monthly else "", line.worked_days, line.paid_days, line.base_salary if monthly else "",
            line.salary_advance if monthly else "", line.leave_deduction if monthly else "", line.penalty if monthly else "",
            line.allowance if monthly else "", line.gross_wages, line.esi_percentage, line.esi_deduction,
            line.esi_employer_percentage, line.esi_employer_contribution, line.pf_percentage, line.pf_deduction,
            line.pf_employer_percentage, line.pf_employer_contribution, line.overtime_hours, line.effective_overtime_rate,
            line.overtime_amount, line.payable, line.per_hour_salary, line.per_day_salary,
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index, column, float(value) if hasattr(value, "as_tuple") else value)
            cell.font = Font(name="Arial", size=9, color=navy)
            cell.alignment = Alignment(horizontal="left" if column in (2, 3, 4) else "center", vertical="center", wrap_text=column == 3)
            cell.border = Border(bottom=thin)
            if column >= 12 and column not in (18, 20, 22, 24, 26):
                cell.number_format = '#,##0.00;[Red](#,##0.00);-'
        if monthly:
            sheet.cell(index, 7, f"=MAX(0,E{index}-F{index})")
            sheet.cell(index, 11, f"=F{index}+H{index}+I{index}+J{index}")
            sheet.cell(index, 14, f"=ROUND(AE{index}*G{index},2)")
            sheet.cell(index, 17, f"=ROUND(L{index}-N{index}-M{index}-O{index}+P{index},2)")
            sheet.cell(index, 19, f"=ROUND(Q{index}*R{index}/100,2)")
            sheet.cell(index, 21, f"=ROUND(Q{index}*T{index}/100,2)")
            sheet.cell(index, 23, f"=ROUND(Q{index}*V{index}/100,2)")
            sheet.cell(index, 25, f"=ROUND(Q{index}*X{index}/100,2)")
            sheet.cell(index, 28, f"=ROUND(Z{index}*AA{index},2)")
            sheet.cell(index, 29, f"=ROUND(Q{index}-S{index}-W{index}+AB{index},2)")
            sheet.cell(index, 30, f"=ROUND(AE{index}/9,2)")
            sheet.cell(index, 31, f"=ROUND(L{index}/{days},2)")
        else:
            sheet.cell(index, 17, f"=ROUND(J{index}*AE{index}+AB{index},2)")
            sheet.cell(index, 19, f"=ROUND(Q{index}*R{index}/100,2)")
            sheet.cell(index, 21, f"=ROUND(Q{index}*T{index}/100,2)")
            sheet.cell(index, 23, f"=ROUND(Q{index}*V{index}/100,2)")
            sheet.cell(index, 25, f"=ROUND(Q{index}*X{index}/100,2)")
            sheet.cell(index, 28, f"=ROUND(Z{index}*AA{index},2)")
            sheet.cell(index, 29, f"=ROUND(Q{index}-S{index}-W{index},2)")
            sheet.cell(index, 30, f"=ROUND(AE{index}/9,2)")
        sheet.row_dimensions[index].height = 25

    total_row = start + len(payroll.lines)
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    sheet.cell(total_row, 1, f"TOTAL ({len(payroll.lines)} employees)")
    for column in range(1, last_column + 1):
        cell = sheet.cell(total_row, column)
        cell.fill = PatternFill("solid", fgColor="FFEDD5")
        cell.font = Font(name="Arial", size=9, bold=True, color=navy)
        cell.border = Border(top=Side(style="medium", color=orange))
    for column in (10, 11, 13, 14, 15, 16, 17, 19, 21, 23, 25, 26, 28, 29):
        letter = get_column_letter(column)
        sheet.cell(total_row, column, f"=SUM({letter}{start}:{letter}{total_row - 1})")
        sheet.cell(total_row, column).number_format = '#,##0.00;[Red](#,##0.00);-'

    widths = {1: 8, 2: 14, 3: 24, 4: 14}
    for column in range(1, last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = widths.get(column, 13)
    sheet.freeze_panes = "E9"
    sheet.auto_filter.ref = f"A8:{last_letter}{total_row - 1}"
    sheet.print_title_rows = "1:8"
    sheet.print_area = f"A1:{last_letter}{total_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.left.text = f"{supplier.name} | Payroll {payroll.month}"
    sheet.oddFooter.right.text = "Page &P of &N"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_payroll_xlsx(payroll: Payroll) -> bytes:
    """Fill the Artifact Tool template, preserving its styles and embedded logo.

    Each formula has a saved-value cache so Excel and other viewers display the
    same amounts immediately, before their first recalculation.
    """
    if any(line.calculation_method == "Monthly" for line in payroll.lines):
        return export_monthly_register_xlsx(payroll)
    with ZipFile(ASSETS / "payroll-template.xlsx") as source:
        sheet = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
        data = sheet.find(f"{{{NS}}}sheetData")
        prototype = next(row for row in data if row.get("r") == "16")
        count = len(payroll.lines)
        if not count:
            raise ValueError("A payroll register must include employees")
        delta = count - 1
        last, total = 15 + count, 16 + count

        def shifted(ref):
            return re.sub(r"(\$?[A-Z]+\$?)(\d+)", lambda match: match[1] + str(int(match[2]) + (delta if int(match[2]) >= 17 else 0)), ref)

        for row in list(data):
            number = int(row.get("r"))
            if number >= 17:
                row.set("r", str(number + delta))
                for cell in row:
                    cell.set("r", shifted(cell.get("r")))
        index = list(data).index(prototype)
        data.remove(prototype)
        for offset in range(count):
            row = deepcopy(prototype)
            number = 16 + offset
            row.set("r", str(number))
            # Grow wrapped employee names and IDs instead of truncating them.
            person = payroll.lines[offset]
            row.set("ht", str(max(28, 14 * max((len(person.full_name) + 24) // 25, (len(person.employee_id or "") + 11) // 12))))
            for cell in row:
                cell.set("r", re.sub(r"\d+$", str(number), cell.get("r")))
            data.insert(index + offset, row)
        cells = {cell.get("r"): cell for cell in sheet.iter(f"{{{NS}}}c")}

        def write(ref, value, formula=None):
            cell = cells[ref]
            for child in list(cell):
                cell.remove(child)
            cell.attrib.pop("t", None)
            if isinstance(value, str):
                # IDs and names are always text, including leading zeros or '='.
                cell.set("t", "inlineStr")
                text = ET.SubElement(ET.SubElement(cell, f"{{{NS}}}is"), f"{{{NS}}}t")
                text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                text.text = value
            else:
                if formula:
                    ET.SubElement(cell, f"{{{NS}}}f").text = formula
                ET.SubElement(cell, f"{{{NS}}}v").text = str(value)

        supplier, totals = payroll.supplier, payroll.totals
        write("C2", supplier.name)
        write("C3", supplier.address)
        write("C5", " | ".join(filter(None, (supplier.phone, supplier.email))))
        write("C6", " | ".join(filter(None, (supplier.website, f"GSTIN: {supplier.gstin}" if supplier.gstin else ""))))
        write("A9", f"Company: {payroll.company_name}")
        write("G9", f"Payroll month: {month_label(payroll.month)}")
        write("A10", f"Worksite: {payroll.company_location}" if payroll.company_location else "")
        write("G10", f"Employees: {count}")
        # Long company names remain visible in the merged header.
        for number, length, chars in ((9, len(payroll.company_name) + 9, 85), (10, len(payroll.company_location) + 10, 85)):
            next(row for row in data if row.get("r") == str(number)).set("ht", str(max(22, 14 * ((length + chars - 1) // chars))))
        for offset, line in enumerate(payroll.lines, 16):
            for col, value in {"A": offset - 15, "B": line.employee_id or "", "C": line.full_name + (f" | OT {line.overtime_hours.normalize():g} h x INR {line.overtime_rate}" if line.overtime_hours else ""), "D": line.worked_days, "E": line.daily_rate, "G": line.esi_percentage, "I": line.pf_percentage}.items():
                write(f"{col}{offset}", value)
            for col, value, formula in (
                ("F", line.gross_wages, (f"ROUND(D{offset}*E{offset}+{line.overtime_hours * line.overtime_rate},2)" if line.overtime_hours else f"ROUND(D{offset}*E{offset},2)")),
                ("H", line.esi_deduction, f"ROUND(F{offset}*G{offset}/100,2)"),
                ("J", line.pf_deduction, f"ROUND(F{offset}*I{offset}/100,2)"),
                ("K", line.payable, f"F{offset}-H{offset}-J{offset}"),
            ):
                write(f"{col}{offset}", value, formula)
        write(f"A{total}", f"TOTAL ({count} employees)")
        for col, value in {"D": totals.worked_days, "F": totals.gross_wages, "H": totals.esi_deduction, "J": totals.pf_deduction, "K": totals.payable}.items():
            write(f"{col}{total}", value, f"SUM({col}16:{col}{last})")
        for col, target, value in (("A", "F", totals.gross_wages), ("D", "H", totals.esi_deduction), ("G", "J", totals.pf_deduction), ("J", "K", totals.payable)):
            write(f"{col}13", value, f"{target}{total}")
        notes_row = 19 + delta
        write(f"A{notes_row}", f"Notes: {payroll.notes}" if payroll.notes else "")
        next(row for row in data if row.get("r") == str(notes_row)).set("ht", str(max(22, 14 * sum(max(1, (len(part) + 139) // 140) for part in payroll.notes.split("\n")))))
        merges = sheet.find(f"{{{NS}}}mergeCells")
        for merge in merges:
            merge.set("ref", shifted(merge.get("ref")))
        dimension = sheet.find(f"{{{NS}}}dimension")
        if dimension is not None:
            dimension.set("ref", f"A1:K{20 + delta}")

        # Insert worksheet children in OOXML order, retaining all drawing links.
        order = ["sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "sheetCalcPr", "sheetProtection", "protectedRanges", "scenarios", "autoFilter", "sortState", "dataConsolidate", "customSheetViews", "mergeCells", "phoneticPr", "conditionalFormatting", "dataValidations", "hyperlinks", "printOptions", "pageMargins", "pageSetup", "headerFooter", "rowBreaks", "colBreaks", "customProperties", "cellWatches", "ignoredErrors", "smartTags", "drawing", "legacyDrawing", "legacyDrawingHF", "picture", "oleObjects", "controls", "webPublishItems", "tableParts", "extLst"]

        def child(name):
            existing = sheet.find(f"{{{NS}}}{name}")
            if existing is not None:
                return existing
            element = ET.Element(f"{{{NS}}}{name}")
            rank = order.index(name)
            position = next((i for i, item in enumerate(sheet) if item.tag.rsplit("}", 1)[-1] in order and order.index(item.tag.rsplit("}", 1)[-1]) > rank), len(sheet))
            sheet.insert(position, element)
            return element

        child("autoFilter").set("ref", f"A15:K{last}")
        validations = child("dataValidations")
        for item in list(validations):
            validations.remove(item)
        days = calendar.monthrange(int(payroll.month[:4]), int(payroll.month[5:]))[1]
        for col, limit, title in (("D", days, "Worked days"), ("E", 1_000_000_000, "Daily rate")):
            rule = ET.SubElement(validations, f"{{{NS}}}dataValidation", type="decimal", operator="between", allowBlank="1", showErrorMessage="1", errorStyle="stop", errorTitle=title, error=f"Enter a number from 0 to {limit}.", sqref=f"{col}16:{col}{last}")
            ET.SubElement(rule, f"{{{NS}}}formula1").text = "0"
            ET.SubElement(rule, f"{{{NS}}}formula2").text = str(limit)
        for col in ("G", "I"):
            rule = ET.SubElement(validations, f"{{{NS}}}dataValidation", type="custom", allowBlank="1", showErrorMessage="1", errorStyle="stop", errorTitle="Deduction percentage", error="Enter percentages from 0 to 100; combined deductions cannot exceed gross wages.", sqref=f"{col}16:{col}{last}")
            ET.SubElement(rule, f"{{{NS}}}formula1").text = "AND($G16>=0,$I16>=0,$G16+$I16<=100,ROUND($F16*$G16/100,2)+ROUND($F16*$I16/100,2)<=$F16)"
        validations.set("count", "4")
        properties = child("sheetPr")
        setup = properties.find(f"{{{NS}}}pageSetUpPr")
        if setup is None:
            setup = ET.SubElement(properties, f"{{{NS}}}pageSetUpPr")
        setup.set("fitToPage", "1")
        child("pageSetup").attrib.update(paperSize="9", orientation="landscape", fitToWidth="1", fitToHeight="0")
        child("pageMargins").attrib.update(left="0.25", right="0.25", top="0.35", bottom="0.35", header="0.15", footer="0.15")
        footer = child("headerFooter")
        odd = footer.find(f"{{{NS}}}oddFooter")
        if odd is None:
            odd = ET.SubElement(footer, f"{{{NS}}}oddFooter")
        odd.text = f"&L{supplier.name.replace('&', '&&')} | {payroll.month}&RPage &P of &N"
        workbook = ET.fromstring(source.read("xl/workbook.xml"))
        defined = workbook.find(f"{{{NS}}}definedNames")
        if defined is None:
            defined = ET.Element(f"{{{NS}}}definedNames")
            calc = workbook.find(f"{{{NS}}}calcPr")
            workbook.insert(list(workbook).index(calc) if calc is not None else len(workbook), defined)
        for item in list(defined):
            if item.get("name") in ("_xlnm.Print_Area", "_xlnm.Print_Titles"):
                defined.remove(item)
        ET.SubElement(defined, f"{{{NS}}}definedName", name="_xlnm.Print_Area", localSheetId="0").text = f"'Payroll'!$A$1:$K${20 + delta}"
        ET.SubElement(defined, f"{{{NS}}}definedName", name="_xlnm.Print_Titles", localSheetId="0").text = "'Payroll'!$1:$15"
        calc = workbook.find(f"{{{NS}}}calcPr")
        if calc is None:
            calc = ET.SubElement(workbook, f"{{{NS}}}calcPr")
        calc.attrib.update(calcMode="auto", fullCalcOnLoad="1")
        output = BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as target:
            for name in source.namelist():
                contents = source.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    contents = ET.tostring(sheet, encoding="utf-8", xml_declaration=True)
                elif name == "xl/workbook.xml":
                    contents = ET.tostring(workbook, encoding="utf-8", xml_declaration=True)
                target.writestr(name, contents)
    return output.getvalue()


def export_monthly_register_pdf(payroll: Payroll) -> bytes:
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("PayrollSans", "Vera.ttf"), ("PayrollBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    navy, orange, muted = (colors.HexColor(value) for value in ("#0F2444", "#C2410C", "#64748B"))
    output = BytesIO()
    page = landscape(A3)
    doc = SimpleDocTemplate(output, pagesize=page, leftMargin=20, rightMargin=20, topMargin=28, bottomMargin=28,
                            title=f"Salary register - {payroll.company_name} - {payroll.month}", author=payroll.supplier.name)
    body = ParagraphStyle("RegisterBody", fontName="PayrollSans", fontSize=5.5, leading=7, textColor=navy, alignment=TA_LEFT)
    center = ParagraphStyle("RegisterCenter", parent=body, alignment=TA_CENTER)
    bold = ParagraphStyle("RegisterBold", parent=body, fontName="PayrollBold")
    heading = ParagraphStyle("RegisterHeading", parent=bold, fontSize=13, leading=16)
    tablehead = ParagraphStyle("RegisterHead", parent=center, fontName="PayrollBold", textColor=colors.white)

    def p(value, style=body):
        return Paragraph(escape(str(value)).replace("\n", "<br/>"), style)

    def money(value):
        return f"{value:,.2f}"

    supplier, totals = payroll.supplier, payroll.totals
    with ZipFile(ASSETS / "invoice-template.xlsx") as archive:
        logo = Image(BytesIO(archive.read("xl/media/image.png")), width=82, height=56)
    brand = [p(supplier.name, heading), p(supplier.address), p(" | ".join(filter(None, (supplier.phone, supplier.email, supplier.website))), body)]
    header = Table([[logo, brand]], colWidths=[95, doc.width - 95])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story = [header, p("SALARY REGISTER", heading), p(f"{payroll.company_name} | {month_label(payroll.month)} | {payroll.company_location or 'Worksite not supplied'}", bold), Spacer(1, 8)]
    summary = Table([[p(label, body) for label in ("Employees", "Gross salary (INR)", "Employee ESI + PF (INR)", "Employer ESI + PF (INR)", "Net payable (INR)")],
                     [p(value, bold) for value in (totals.employee_count, money(totals.gross_wages), money(totals.esi_deduction + totals.pf_deduction), money(totals.esi_employer_contribution + totals.pf_employer_contribution), money(totals.payable))]],
                    colWidths=[doc.width / 5] * 5)
    summary.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")), ("LINEABOVE", (0, 0), (-1, 0), .6, orange), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    story.extend([summary, Spacer(1, 8)])
    labels = ("No.", "Employee ID", "Name", "Joined", "Absent", "Credit", "LOP", "Govt", "Sun", "Present", "Paid", "Salary", "Loan / advance", "Leave deduction", "Penalty", "Allowance", "Gross", "Emp ESI", "Er ESI", "Emp PF", "Er PF", "OT amount", "Net")
    rows = [[p(label, tablehead) for label in labels]]
    for number, line in enumerate(payroll.lines, 1):
        rows.append([p(number, center), p(line.employee_id or "—"), p(line.full_name), p(line.start_date or "—"),
                     p(f"{line.days_absent:g}", center), p(f"{line.credit_leave:g}", center), p(f"{line.loss_of_pay:g}", center),
                     p(f"{line.government_holidays:g}", center), p(line.sunday_holidays, center), p(f"{line.worked_days:g}", center),
                     p(f"{line.paid_days:g}", center), p(money(line.base_salary)), p(money(line.salary_advance)), p(money(line.leave_deduction)),
                     p(money(line.penalty)), p(money(line.allowance)), p(money(line.gross_wages)), p(money(line.esi_deduction)),
                     p(money(line.esi_employer_contribution)), p(money(line.pf_deduction)), p(money(line.pf_employer_contribution)),
                     p(money(line.overtime_amount)), p(money(line.payable), bold)])
    rows.append([p("TOTAL", bold), "", "", "", "", "", "", "", "", p(f"{totals.worked_days:g}", center), p(f"{totals.paid_days:g}", center), "", "", p(money(totals.leave_deduction), bold), "", "", p(money(totals.gross_wages), bold), p(money(totals.esi_deduction), bold), p(money(totals.esi_employer_contribution), bold), p(money(totals.pf_deduction), bold), p(money(totals.pf_employer_contribution), bold), p(money(totals.overtime_amount), bold), p(money(totals.payable), bold)])
    weights = [.025, .055, .09, .05, .035, .035, .03, .03, .025, .035, .035, .055, .06, .06, .045, .045, .06, .05, .05, .05, .05, .05, .06]
    table = LongTable(rows, colWidths=[doc.width * weight / sum(weights) for weight in weights], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), navy), ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
                               ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFEDD5")), ("SPAN", (0, -1), (3, -1)),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 1), (-1, -1), .25, colors.HexColor("#E2E8F0")),
                               ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                               ("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2)]))
    story.extend([table, Spacer(1, 7), p("LOP = absent days - credit leave. Paid days = present + credit leave + government holidays + Sundays. Net = gross - employee ESI - employee PF + overtime amount.", body)])

    def footer(canvas, document):
        canvas.saveState(); canvas.setFont("PayrollSans", 6); canvas.setFillColor(muted)
        canvas.drawString(20, 14, f"{supplier.name[:65]} | Payroll {payroll.month}")
        canvas.drawRightString(page[0] - 20, 14, f"Page {document.page}"); canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_payroll_pdf(payroll: Payroll) -> bytes:
    if any(line.calculation_method == "Monthly" for line in payroll.lines):
        return export_monthly_register_pdf(payroll)
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("PayrollSans", "Vera.ttf"), ("PayrollBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    navy, orange, muted = (colors.HexColor(value) for value in ("#0F2444", "#C2410C", "#64748B"))
    output = BytesIO()
    page = landscape(A4)
    doc = SimpleDocTemplate(output, pagesize=page, leftMargin=28, rightMargin=28, topMargin=34, bottomMargin=34,
                            title=f"Payroll - {payroll.company_name} - {payroll.month}", author=payroll.supplier.name)
    body = ParagraphStyle("PayrollBody", fontName="PayrollSans", fontSize=8, leading=11, textColor=navy, alignment=TA_LEFT)
    small = ParagraphStyle("PayrollSmall", parent=body, fontSize=7, leading=10)
    bold = ParagraphStyle("PayrollStrong", parent=body, fontName="PayrollBold")
    heading = ParagraphStyle("PayrollHeading", parent=bold, fontSize=13, leading=17)
    muted_style = ParagraphStyle("PayrollMuted", parent=small, textColor=muted)
    center = ParagraphStyle("PayrollCenter", parent=body, alignment=TA_CENTER)
    tablehead = ParagraphStyle("PayrollTableHead", parent=small, fontName="PayrollBold", textColor=colors.white)

    def p(text, style=body):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    def money(value):
        return f"{value:,.2f}"

    supplier, totals = payroll.supplier, payroll.totals
    with ZipFile(ASSETS / "invoice-template.xlsx") as archive:
        logo = Image(BytesIO(archive.read("xl/media/image.png")), width=92, height=62.7)
    brand = [p(supplier.name, heading), p(supplier.address), p(" | ".join(filter(None, (supplier.phone, supplier.email))), small),
             p(" | ".join(filter(None, (supplier.website, f"GSTIN: {supplier.gstin}" if supplier.gstin else ""))), small)]
    header = Table([[logo, brand]], colWidths=[110, doc.width - 110])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    story = [header, Spacer(1, 5), p("PAYROLL REGISTER", heading), Spacer(1, 8)]
    details = Table([[[p("COMPANY", muted_style), p(payroll.company_name, bold), p(payroll.company_location, small)],
                      [p("PAYROLL MONTH", muted_style), p(month_label(payroll.month), bold), p(f"{totals.employee_count} employees | {totals.worked_days:g} total worked days", small)]]],
                    colWidths=[doc.width * .62, doc.width * .38])
    details.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F1F5F9")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story.extend([details, Spacer(1, 10)])
    summary = Table([[p(label, small) for label in ("Gross wages (INR)", "ESI deduction (INR)", "PF deduction (INR)", "Payable (INR)")],
                     [p(money(value), heading) for value in (totals.gross_wages, totals.esi_deduction, totals.pf_deduction, totals.payable)]],
                    colWidths=[doc.width / 4] * 4)
    summary.setStyle(TableStyle([("BACKGROUND", (0, 0), (2, -1), colors.HexColor("#F8FAFC")), ("BACKGROUND", (3, 0), (3, -1), colors.HexColor("#FFEDD5")), ("LINEABOVE", (0, 0), (-1, 0), .6, orange), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    story.extend([summary, Spacer(1, 12)])
    labels = ("No.", "Employee ID", "Employee name", "Worked days", "Daily rate (INR)", "Gross wages (INR)", "ESI %", "ESI deduction (INR)", "PF %", "PF deduction (INR)", "Payable (INR)")
    rows = [[p(label, tablehead) for label in labels]]
    for number, line in enumerate(payroll.lines, 1):
        rows.append([p(number, center), p(line.employee_id or "—"), p(line.full_name + (f"\nOT: {line.overtime_hours.normalize():g} h x INR {money(line.overtime_rate)}" if line.overtime_hours else "")), p(f"{line.worked_days:g}", center),
                     p(money(line.daily_rate)), p(money(line.gross_wages)), p(f"{line.esi_percentage:g}", center),
                     p(money(line.esi_deduction)), p(f"{line.pf_percentage:g}", center), p(money(line.pf_deduction)), p(money(line.payable), bold)])
    rows.append([p("TOTAL", bold), "", "", p(f"{totals.worked_days:g}", center), "", p(money(totals.gross_wages), bold), "", p(money(totals.esi_deduction), bold), "", p(money(totals.pf_deduction), bold), p(money(totals.payable), bold)])
    widths = [.035, .09, .20, .07, .09, .105, .055, .10, .055, .10, .10]
    table = LongTable(rows, colWidths=[doc.width * width for width in widths], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), navy), ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
                               ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#FFEDD5")), ("SPAN", (0, -1), (2, -1)),
                               ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 1), (-1, -1), .3, colors.HexColor("#E2E8F0")),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                               ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]))
    story.extend([table, Spacer(1, 10), p("Gross wages = worked days x daily rate + overtime hours x overtime hourly rate. ESI and PF each use the entered percentage of gross wages. Payable = gross wages - ESI - PF.", muted_style)])
    if payroll.notes:
        story.append(KeepTogether([Spacer(1, 7), p("Notes", bold), p(payroll.notes)]))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("PayrollSans", 7)
        canvas.setFillColor(muted)
        canvas.drawString(28, 18, f"{supplier.name[:65]} | Payroll {payroll.month}")
        canvas.drawRightString(page[0] - 28, 18, f"Page {document.page}")
        if document.page > 1:
            canvas.drawString(28, page[1] - 20, f"PAYROLL REGISTER | {payroll.company_name[:80]} | {month_label(payroll.month)}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
